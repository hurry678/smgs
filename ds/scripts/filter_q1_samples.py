#!/usr/bin/env python3
"""Filter and audit the 100 raw Q1 samples.

This script only performs data-quality filtering. It never modifies source
videos or labels. It writes machine-readable qualified/excluded manifests.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import subprocess
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ALLOWED_ANNOTATIONS = {"Negative", "Neutral", "Positive"}
REQUIRED_COLUMNS = ["video_id", "clip_id", "text", "label", "annotation"]


def normalize_id(value: Any) -> str:
    """Normalize Excel IDs without losing a leading minus sign."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def run_command(command: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        command,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )


def tool_version(executable: str) -> str:
    result = run_command([executable, "-version"])
    if result.returncode != 0:
        return "unavailable"
    return result.stdout.splitlines()[0] if result.stdout else "unknown"


def probe_media(path: Path) -> tuple[bool, dict[str, Any], str]:
    command = [
        "ffprobe",
        "-v",
        "error",
        "-show_entries",
        "format=duration,format_name:stream=index,codec_type,codec_name,width,height,r_frame_rate,sample_rate,channels",
        "-of",
        "json",
        str(path),
    ]
    result = run_command(command)
    if result.returncode != 0:
        return False, {}, result.stderr.strip() or "ffprobe failed"
    try:
        return True, json.loads(result.stdout or "{}"), ""
    except json.JSONDecodeError as exc:
        return False, {}, f"invalid ffprobe JSON: {exc}"


def parse_float(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def decode_media(path: Path) -> tuple[bool, str]:
    """Fully decode the first video and audio streams to catch corrupt media."""
    command = [
        "ffmpeg",
        "-nostdin",
        "-v",
        "error",
        "-xerror",
        "-i",
        str(path),
        "-map",
        "0:v:0",
        "-map",
        "0:a:0",
        "-f",
        "null",
        "-",
    ]
    result = run_command(command)
    if result.returncode == 0:
        return True, ""
    return False, result.stderr.strip() or "ffmpeg decode failed"


def load_label_rows(excel_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    if "label" not in workbook.sheetnames:
        raise ValueError(f"Worksheet 'label' not found. Sheets: {workbook.sheetnames}")
    worksheet = workbook["label"]
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as exc:
        raise ValueError("The label worksheet is empty") from exc
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")
    result: list[dict[str, Any]] = []
    for values in rows:
        if not any(value is not None and str(value).strip() for value in values):
            continue
        record = dict(zip(headers, values))
        result.append(record)
    workbook.close()
    return result


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def build_audit_row(
    record: dict[str, Any],
    video_root: Path,
    duplicate_ids: set[str],
    run_decode: bool,
) -> dict[str, Any]:
    video_id = normalize_id(record.get("video_id"))
    clip_id = normalize_id(record.get("clip_id"))
    sample_id = f"{video_id}$_${clip_id}"
    text = "" if record.get("text") is None else str(record.get("text")).strip()
    label_value = parse_float(record.get("label"))
    annotation = "" if record.get("annotation") is None else str(record.get("annotation")).strip()
    video_path = video_root / video_id / f"{clip_id}.mp4"

    reasons: list[str] = []
    if not video_id:
        reasons.append("missing_video_id")
    if not clip_id:
        reasons.append("missing_clip_id")
    if sample_id in duplicate_ids:
        reasons.append("duplicate_sample_id")
    if not text:
        reasons.append("missing_text")
    if label_value is None:
        reasons.append("missing_or_invalid_label")
    elif not (-3.0 <= label_value <= 3.0):
        reasons.append("label_out_of_range")
    if annotation not in ALLOWED_ANNOTATIONS:
        reasons.append("missing_or_invalid_annotation")

    file_exists = video_path.is_file()
    file_size = video_path.stat().st_size if file_exists else 0
    if not file_exists:
        reasons.append("missing_video_file")
    elif file_size <= 0:
        reasons.append("empty_video_file")

    probe_ok = False
    decode_ok = False
    duration = None
    video_codec = audio_codec = ""
    width = height = None
    frame_rate = None
    sample_rate = channels = None
    video_stream_count = audio_stream_count = 0
    probe_error = ""
    decode_error = ""

    if file_exists and file_size > 0:
        probe_ok, probe_data, probe_error = probe_media(video_path)
        if not probe_ok:
            reasons.append("ffprobe_failed")
        else:
            streams = probe_data.get("streams", [])
            video_streams = [stream for stream in streams if stream.get("codec_type") == "video"]
            audio_streams = [stream for stream in streams if stream.get("codec_type") == "audio"]
            video_stream_count = len(video_streams)
            audio_stream_count = len(audio_streams)
            if video_stream_count == 0:
                reasons.append("missing_video_stream")
            if audio_stream_count == 0:
                reasons.append("missing_audio_stream")
            duration = parse_float(probe_data.get("format", {}).get("duration"))
            if duration is None or duration <= 0:
                reasons.append("missing_or_invalid_duration")
            if video_streams:
                stream = video_streams[0]
                video_codec = str(stream.get("codec_name") or "")
                width = stream.get("width")
                height = stream.get("height")
                frame_rate = str(stream.get("r_frame_rate") or "")
            if audio_streams:
                stream = audio_streams[0]
                audio_codec = str(stream.get("codec_name") or "")
                sample_rate = stream.get("sample_rate")
                channels = stream.get("channels")
        if run_decode and probe_ok:
            decode_ok, decode_error = decode_media(video_path)
            if not decode_ok:
                reasons.append("full_decode_failed")

    annotation_consistent = None
    if label_value is not None and annotation in ALLOWED_ANNOTATIONS:
        expected = "Negative" if label_value < 0 else "Positive" if label_value > 0 else "Neutral"
        annotation_consistent = annotation == expected
        if not annotation_consistent:
            reasons.append("annotation_label_mismatch")

    reasons = list(dict.fromkeys(reasons))
    return {
        "sample_id": sample_id,
        "video_id": video_id,
        "clip_id": clip_id,
        "text": text,
        "text_chars": len(text),
        "label": label_value,
        "annotation": annotation,
        "source_video_path": str(video_path),
        "file_exists": file_exists,
        "file_size_bytes": file_size,
        "ffprobe_ok": probe_ok,
        "video_stream_count": video_stream_count,
        "audio_stream_count": audio_stream_count,
        "video_codec": video_codec,
        "width": width,
        "height": height,
        "frame_rate": frame_rate,
        "audio_codec": audio_codec,
        "sample_rate": sample_rate,
        "channels": channels,
        "duration_sec": duration,
        "full_decode_ok": decode_ok,
        "text_nonempty": bool(text),
        "label_valid": label_value is not None and -3.0 <= label_value <= 3.0,
        "annotation_valid": annotation in ALLOWED_ANNOTATIONS,
        "annotation_consistent": annotation_consistent,
        "quality_status": "excluded" if reasons else "qualified",
        "exclusion_reasons": ";".join(reasons),
        "probe_error": probe_error,
        "decode_error": decode_error,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Filter Q1 raw samples by basic data quality.")
    parser.add_argument(
        "--excel",
        type=Path,
        default=Path(r"C:\work\数模\中文题目\E题\E题数据\附件1-数据集原始多模态样本\MOSEI数据集部分原始视频-100条\label-100.xlsx"),
    )
    parser.add_argument(
        "--video-root",
        type=Path,
        default=Path(r"C:\work\数模\中文题目\E题\E题数据\附件1-数据集原始多模态样本\MOSEI数据集部分原始视频-100条"),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(r"C:\work\数模\中文题目\E题\artifacts\q1\data_cleaning"),
    )
    parser.add_argument("--skip-full-decode", action="store_true")
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    records = load_label_rows(args.excel)
    sample_ids = [f"{normalize_id(row.get('video_id'))}$_${normalize_id(row.get('clip_id'))}" for row in records]
    duplicate_ids = {sample_id for sample_id, count in Counter(sample_ids).items() if count > 1}

    rows = [
        build_audit_row(record, args.video_root, duplicate_ids, not args.skip_full_decode)
        for record in records
    ]
    qualified = [row for row in rows if row["quality_status"] == "qualified"]
    excluded = [row for row in rows if row["quality_status"] == "excluded"]

    fieldnames = list(rows[0].keys()) if rows else []
    write_csv(args.output_dir / "all_samples_quality.csv", rows, fieldnames)
    write_csv(args.output_dir / "qualified_samples.csv", qualified, fieldnames)
    write_csv(args.output_dir / "excluded_samples.csv", excluded, fieldnames)

    reason_counts = Counter(
        reason
        for row in excluded
        for reason in row["exclusion_reasons"].split(";")
        if reason
    )
    report = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "scope": "问题1附件1原始样本基础数据清洗；不做特征提取、不做训练",
        "inputs": {
            "excel": str(args.excel),
            "video_root": str(args.video_root),
        },
        "tools": {
            "ffprobe": tool_version("ffprobe"),
            "ffmpeg": tool_version("ffmpeg"),
            "openpyxl": __import__("openpyxl").__version__,
        },
        "rules": {
            "required_columns": REQUIRED_COLUMNS,
            "allowed_annotations": sorted(ALLOWED_ANNOTATIONS),
            "text": "去除首尾空白后必须非空",
            "media": "视频文件必须存在且非空；至少一条视频流和一条音频流；时长必须为正；默认执行完整解码",
            "label": "数值必须在[-3,3]且annotation必须与符号一致",
            "duplicates": "重复sample_id全部标记为不合格",
            "source_data_policy": "只生成清单和报告，不复制、不删除、不修改原始样本",
        },
        "counts": {
            "excel_rows": len(records),
            "audited": len(rows),
            "qualified": len(qualified),
            "excluded": len(excluded),
            "duplicate_sample_ids": len(duplicate_ids),
            "full_decode_checked": 0 if args.skip_full_decode else len(rows),
        },
        "exclusion_reason_counts": dict(sorted(reason_counts.items())),
        "outputs": {
            "all_samples_quality": str(args.output_dir / "all_samples_quality.csv"),
            "qualified_samples": str(args.output_dir / "qualified_samples.csv"),
            "excluded_samples": str(args.output_dir / "excluded_samples.csv"),
        },
    }
    (args.output_dir / "quality_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report["counts"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
