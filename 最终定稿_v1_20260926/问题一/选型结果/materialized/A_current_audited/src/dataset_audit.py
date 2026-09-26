"""Read-only audit of all supplied attachments and Q1 outputs; no fitting or cleaning."""
import csv
import gc
import hashlib
import itertools
import json
import math
import re
import time
from collections import Counter, defaultdict
import numpy as np
from scipy.stats import pearsonr, spearmanr
from prepare import ROOT, dump, sha256
from check_text import ArrayUnpickler

DATA = ROOT.parent / "E题数据"
OUT = ROOT / "reports/data_quality"
FEATURE_NAMES = ([f"bert_{i:03d}" for i in range(128)] + [f"mfcc_{i:02d}" for i in range(13)]
                 + ["log_rms", "zcr", "f0_hz", "eye_a", "eye_b", "mouth_open", "mouth_width",
                    "corner_61", "corner_291"])


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fields = list(dict.fromkeys(k for row in rows for k in row))
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def statistics(values):
    x = np.asarray(values, dtype=float).ravel()
    valid = x[np.isfinite(x)]
    result = {"n_total": len(x), "n_valid": len(valid), "n_missing": int((~np.isfinite(x)).sum())}
    if not len(valid):
        return {**result, **{k: None for k in ("mean", "std", "min", "q25", "median", "q75",
                                              "p90", "max", "iqr_lower", "iqr_upper", "iqr_outliers")}}
    quantiles = np.quantile(valid, [0, .25, .5, .75, .9, 1])
    lo, hi = quantiles[1] - 1.5*(quantiles[3]-quantiles[1]), quantiles[3] + 1.5*(quantiles[3]-quantiles[1])
    return {**result, "mean": float(valid.mean()), "std": float(valid.std(ddof=1)) if len(valid)>1 else None,
            **dict(zip(("min", "q25", "median", "q75", "p90", "max"), map(float, quantiles))),
            "iqr_lower": float(lo), "iqr_upper": float(hi),
            "iqr_outliers": int(((valid < lo) | (valid > hi)).sum())}


def correlation_table(rows, columns):
    """Pairwise finite observations, one row per sample; no significance/causal claims."""
    result = []
    for a, b in itertools.combinations(columns, 2):
        x, y = np.array([[r.get(a, np.nan), r.get(b, np.nan)] for r in rows], float).T
        finite = np.isfinite(x) & np.isfinite(y)
        x, y = x[finite], y[finite]
        computable = len(x) >= 3 and np.ptp(x) > 0 and np.ptp(y) > 0
        result.append({"x": a, "y": b, "n": len(x),
                       "pearson_r": float(pearsonr(x, y).statistic) if computable else None,
                       "spearman_rho": float(spearmanr(x, y).statistic) if computable else None})
    return result


def duplicates(keys, ids):
    grouped = defaultdict(list)
    for key, sid in zip(keys, ids):
        grouped[str(key)].append(sid)
    return [v for v in grouped.values() if len(v) > 1]


def workbook_audit(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheet = book["label"]
    cells = list(sheet.iter_rows())
    headers = [c.value for c in cells[0]]
    records, missing, formulas, errors = [], Counter(), [], []
    for line, cells_row in enumerate(cells[1:], 2):
        if not any(c.value is not None for c in cells_row):
            continue
        row = dict(zip(headers, [c.value for c in cells_row]))
        sid = str(row.get("video_id")) + "$_$" + str(row.get("clip_id"))
        for name, cell in zip(headers, cells_row):
            if cell.value is None or not str(cell.value).strip():
                missing[name] += 1
            if cell.data_type == "f":
                formulas.append({"row": line, "column": name})
        try:
            value = float(row["label"])
            expected = "Neutral" if value == 0 else "Positive" if value > 0 else "Negative"
            if not math.isfinite(value) or not -3 <= value <= 3 or row["annotation"] != expected:
                errors.append({"id": sid, "error": "label_range_or_sign_annotation"})
        except (TypeError, ValueError):
            errors.append({"id": sid, "error": "label_not_numeric"})
        if not re.fullmatch(r"[A-Za-z0-9_-]+", str(row.get("video_id", ""))):
            errors.append({"id": sid, "error": "video_id_format"})
        if not re.fullmatch(r"\d+", str(row.get("clip_id", ""))):
            errors.append({"id": sid, "error": "clip_id_format"})
        if not isinstance(row.get("text"), str) or not row["text"].strip():
            errors.append({"id": sid, "error": "empty_or_nonstring_text"})
        elif any(ord(c) < 32 and c not in "\t\r\n" for c in row["text"]):
            errors.append({"id": sid, "error": "text_control_characters"})
        records.append({**row, "sample_id": sid})
    # Do not export the spreadsheet's historical private absolute source path.
    historical_note = {}
    if "修复说明" in book.sheetnames:
        for key, value, *_ in book["修复说明"].values:
            if key in {"筛选后数据行数", "列数", "负向样本", "中性样本", "正向样本", "训练集", "验证集", "测试集"}:
                historical_note[key] = value
    book.close()
    ids = [r["sample_id"] for r in records]
    return records, {
        "file": str(path.relative_to(DATA)), "sha256": sha256(path), "rows": len(records), "columns": headers,
        "duplicate_headers": [k for k, count in Counter(headers).items() if count > 1],
        "missing_cells_by_column": dict(missing), "formula_cells": formulas, "format_or_logic_errors": errors,
        "duplicate_ids": duplicates(ids, ids),
        "numeric_video_id_rows": sum(str(r["video_id"]).isdigit() for r in records),
        "duplicate_full_rows": duplicates([json.dumps(r, sort_keys=True) for r in records], ids),
        "duplicate_normalized_text": duplicates([" ".join(str(r["text"]).lower().split()) for r in records], ids),
        "historical_note_unverified": historical_note,
        "numeric_labels_stored_as_text": sum(isinstance(r["label"], str) for r in records),
    }


def probe_original(row, path):
    """Decode all original frames/channels. Threshold flags describe observations only."""
    import av
    audio_parts, pts, video_frames = [], [], 0
    thumbnail_hashes, luminance = [], []
    next_image = 0.0
    with av.open(str(path)) as container:
        audio = container.streams.audio[0]
        rate, channel_count = audio.codec_context.sample_rate, audio.codec_context.channels
        for frame in container.decode(audio):
            x = frame.to_ndarray()
            if not frame.format.is_planar:
                x = x.reshape(-1, channel_count).T
            x = x.astype(np.float64) / (np.iinfo(x.dtype).max + 1) if np.issubdtype(x.dtype, np.integer) else x
            audio_parts.append(x)
    sound = np.concatenate(audio_parts, axis=1)
    with av.open(str(path)) as container:
        for frame in container.decode(video=0):
            t = float(frame.pts * frame.time_base)
            pts.append(t)
            video_frames += 1
            if t >= next_image:
                image = frame.reformat(width=64, height=36, format="gray").to_ndarray()
                thumbnail_hashes.append(hashlib.sha256(image.tobytes()).hexdigest())
                luminance.append((float(image.mean()), float(image.std())))
                next_image = t + 1.0
    with (ROOT / "intermediate/probes" / (row["sample_id"] + ".json")).open() as handle:
        old = json.load(handle)
    saved_pts = np.array([float(f["best_effort_timestamp_time"]) for f in old["frames"]])
    correct_pts = len(pts) == len(saved_pts) and np.allclose(pts, saved_pts, atol=5.1e-7, rtol=0)
    rms = float(np.sqrt(np.mean(sound.astype(np.float64)**2)))
    return {
        "sample_id": row["sample_id"], "decoded_frames": video_frames, "decoded_audio_channels": channel_count,
        "audio_sample_rate": rate, "decoded_audio_samples_per_channel": sound.shape[1],
        "audio_duration_s": sound.shape[1] / rate, "audio_nonfinite": int((~np.isfinite(sound)).sum()),
        "audio_nonzero": int(np.count_nonzero(sound)), "digital_silence": bool(not np.any(sound)),
        "audio_rms_dbfs": float(20*np.log10(rms)) if rms else None,
        "audio_peak_abs": float(np.max(np.abs(sound))),
        "near_full_scale_fraction": float((np.abs(sound) >= .999).mean()),
        "pts_strictly_increasing": bool(len(pts) > 0 and np.all(np.diff(pts) > 0)),
        "original_pts_match_saved_probe": bool(correct_pts),
        "black_thumbnail_candidates": sum(m < 5 and s < 5 for m, s in luminance),
        "sampled_thumbnails": len(luminance),
        "repeated_adjacent_thumbnails": sum(a == b for a, b in zip(thumbnail_hashes, thumbnail_hashes[1:])),
    }


def q1_audit(decode=True):
    rows = json.loads((ROOT / "resources/manifest.json").read_text())
    label_path = next(DATA.glob("附件1*/**/label-100.xlsx"))
    labels, sheet = workbook_audit(label_path)
    label_map = {r["sample_id"]: r for r in labels}
    input_hashes = {r["path"]: r["sha256"] for r in json.loads((ROOT / "resources/input_manifest.json").read_text())}
    actual = list(label_path.parent.rglob("*.mp4"))
    ids = [r["sample_id"] for r in rows]
    source_hashes, problems, media, samples, phrase_rows = [], [], [], [], []
    all_x, all_valid = [], []
    for index, row in enumerate(rows):
        sid = row["sample_id"]
        original = label_path.parent / row["video_id"] / (row["clip_id"] + ".mp4")
        original_sha = sha256(original)
        source_hashes.append(original_sha)
        if original_sha != sha256(ROOT / row["source_file"]) or original_sha != input_hashes[row["source_file"]]:
            problems.append({"sample_id": sid, "issue": "source_copy_or_manifest_hash_mismatch"})
        if row["text"] != label_map[sid]["text"]:
            problems.append({"sample_id": sid, "issue": "transcript_mismatch"})
        if decode:
            try:
                media.append(probe_original(row, original))
            except Exception as exc:
                problems.append({"sample_id": sid, "issue": "decode_failure", "error": type(exc).__name__})
        meta = json.loads((ROOT / "outputs/details" / (sid + ".json")).read_text())
        cache = ROOT / "intermediate/samples" / sid
        with np.load(cache / "audio.npz") as au, np.load(cache / "vision.npz") as vi:
            pitch = au["f0"][au["f0_valid"]]
            face_valid = vi["valid"].all(1)
            if (not np.isfinite(au["values"]).all() or
                    np.any((au["values"][:,14] < 0) | (au["values"][:,14] > 1)) or
                    np.any((pitch < 65) | (pitch > 500))):
                problems.append({"sample_id":sid,"issue":"acoustic_physical_range_or_finiteness"})
            v=vi["values"][face_valid]
            if (not np.isfinite(v).all() or np.any(v[:,:4] < 0) or
                    np.any(v[:,:2] > 1) or np.any(v[:,2] > .6) or
                    np.any(v[:,3] > 1) or np.any(np.abs(v[:,4:]) > 1)):
                problems.append({"sample_id":sid,"issue":"valid_face_physical_range_or_finiteness"})
            sample = {
                "sample_id": sid, "video_id": row["video_id"], "duration_s": row["duration_s"],
                "label": float(label_map[sid]["label"]), "annotation": label_map[sid]["annotation"],
                "text_chars": len(row["text"]), "text_words": len(row["text"].split()),
                "words_per_video_second": len(row["text"].split()) / row["duration_s"],
                "token_count": meta["token_count"], "phrase_count": len(meta["phrases"]),
                "alignment_hard_passed": meta["alignment_quality"]["hard_passed"],
                "digital_silence": meta["digitally_silent_audio"],
                "native_log_rms_mean": float(au["values"][:, 13].mean()),
                "native_zcr_mean": float(au["values"][:, 14].mean()),
                "voiced_rate": float(au["f0_valid"].mean()),
                "f0_mean_hz": float(pitch.mean()) if len(pitch) else np.nan,
                "face_valid_rate": float(face_valid.mean()),
                "ambiguous_valid_frames": int((vi["ambiguous"] & face_valid).sum()),
                "crop_candidate_frames": int(vi["crop_fallback_used"].sum()),
                "crop_valid_frames": int((vi["crop_fallback_used"] & face_valid).sum()),
                "geometry_rejected_frames": int(vi["geometry_rejected"].sum()),
            }
            for j, name in enumerate(FEATURE_NAMES[144:]):
                sample[f"native_{name}_mean"] = float(vi["values"][face_valid, j].mean()) if face_valid.any() else np.nan
            for p in meta["phrases"]:
                active = (vi["time"] >= p["start_s"]) & (vi["time"] < p["end_s"])
                if not active.any():
                    reason = "sampling_no_coverage"
                elif not face_valid[active].any():
                    reason = "detected_geometry_rejected" if vi["geometry_rejected"][active].any() else "no_reliable_face_detected"
                else:
                    reason = "observed_with_subject_ambiguity" if (vi["ambiguous"][active] & face_valid[active]).any() else "observed"
                duration = p["end_s"] - p["start_s"]
                phrase_rows.append({
                    "sample_id": sid, "phrase_index": p["phrase_index"], "start_s": p["start_s"],
                    "end_s": p["end_s"], "duration_s": duration, "words": len(p["text"].split()),
                    "words_per_second": len(p["text"].split()) / duration, "text": p["text"],
                    "sampled_frames": int(active.sum()), "valid_face_frames": int(face_valid[active].sum()),
                    "visual_observation_state": reason, "flags": ";".join(p["flags"]),
                })
        with np.load(ROOT / meta["feature_file"], allow_pickle=False) as z:
            x, valid = z["features"], z["valid"]
            all_x.append(x)
            all_valid.append(valid)
            durations = np.diff(z["intervals"], axis=1).ravel()
            sample["phrase_time_coverage"] = float(durations.sum()/row["duration_s"]) if len(x) else np.nan
            sample["phrase_visual_observed_rate"] = float(valid[:,144:].all(1).mean()) if len(x) else np.nan
            sample["phrase_pitch_observed_rate"] = float(valid[:,143].mean()) if len(x) else np.nan
            if not np.isfinite(x).all() or np.any(x[~valid] != 0):
                problems.append({"sample_id": sid, "issue": "npz_nonfinite_or_invalid_mask"})
        samples.append(sample)
        if (index+1) % 10 == 0:
            print(f"Q1 audit {index+1}/100", flush=True)
    if decode:
        dump(OUT / "media_integrity.json", media)
    else:
        media = json.loads((OUT / "media_integrity.json").read_text())
    media_by_id = {r["sample_id"]: r for r in media}
    for sample in samples:
        obs = media_by_id.get(sample["sample_id"], {})
        sample.update({k: obs.get(k) for k in ("audio_rms_dbfs", "audio_peak_abs", "near_full_scale_fraction",
                                              "black_thumbnail_candidates", "repeated_adjacent_thumbnails")})
    x, mask = np.concatenate(all_x), np.concatenate(all_valid)
    dimensions = [{"dimension": j, "name": name, **statistics(np.where(mask[:,j], x[:,j], np.nan))}
                  for j, name in enumerate(FEATURE_NAMES)]
    numeric = ["duration_s", "label", "text_chars", "text_words", "token_count", "phrase_count",
               "words_per_video_second", "phrase_time_coverage", "face_valid_rate", "voiced_rate", "f0_mean_hz",
               "native_log_rms_mean", "native_zcr_mean"] + [f"native_{n}_mean" for n in FEATURE_NAMES[144:]]
    distribution = [{"metric": k, **statistics([r[k] for r in samples])} for k in numeric]
    candidates = [{"sample_id": r["sample_id"], "metric": d["metric"], "value": r[d["metric"]],
                   "lower": d["iqr_lower"], "upper": d["iqr_upper"]}
                  for d in distribution for r in samples
                  if np.isfinite(r[d["metric"]]) and (
                      r[d["metric"]] < d["iqr_lower"] or r[d["metric"]] > d["iqr_upper"])]
    correlations = correlation_table(samples, numeric)
    # Equal weight per source video as a descriptive dependence sensitivity check.
    groups = defaultdict(list)
    for row in samples:
        groups[row["video_id"]].append(row)
    grouped = []
    for video, clips in groups.items():
        rec = {"video_id": video}
        for key in numeric:
            values = np.array([r[key] for r in clips], float)
            rec[key] = float(values[np.isfinite(values)].mean()) if np.isfinite(values).any() else np.nan
        grouped.append(rec)
    for name, recs in (("q1_samples", samples), ("q1_phrases", phrase_rows), ("q1_feature_dimensions", dimensions),
                       ("q1_distributions", distribution), ("q1_iqr_candidates", candidates),
                       ("q1_correlations", correlations),
                       ("q1_video_group_correlations", correlation_table(grouped, numeric))):
        write_csv(OUT / (name+".csv"), recs)
    actual_ids = {p.parent.name+"$_$"+p.stem for p in actual}
    report = {
        "spreadsheet": sheet, "sample_count": len(rows), "video_groups": len(groups),
        "group_sizes": statistics([len(r) for r in groups.values()]),
        "source_file_count": len(actual), "missing_source_ids": sorted(set(ids)-actual_ids),
        "extra_source_ids": sorted(actual_ids-set(ids)), "source_hash_duplicate_groups": duplicates(source_hashes, ids),
        "original_label_copy_equal": sha256(label_path) == sha256(ROOT / "resources/input/label-100.xlsx"),
        "integrity_errors": problems,
        "label_counts": dict(Counter(r["annotation"] for r in labels)),
        "format_note": "Labels stored as text parse losslessly to numeric; no source cells changed.",
        "duration_outside_stated_range": [r["sample_id"] for r in samples if not 2.648 <= r["duration_s"] <= 34.567],
        "distributions": distribution,
        "digital_silence_ids": [r["sample_id"] for r in media if r["digital_silence"]],
        "decode_checked": len(media),
        "decoded_frame_count": sum(r["decoded_frames"] for r in media),
        "pts_mismatch_ids": [r["sample_id"] for r in media if not r["original_pts_match_saved_probe"]],
        "nonincreasing_pts_ids": [r["sample_id"] for r in media if not r["pts_strictly_increasing"]],
        "near_full_scale_candidates": [r["sample_id"] for r in media if r["near_full_scale_fraction"] > .001],
        "black_thumbnail_candidates": [r["sample_id"] for r in media if r["black_thumbnail_candidates"]],
        "repeated_thumbnail_candidates": [r["sample_id"] for r in media if r["repeated_adjacent_thumbnails"]],
        "phrase_count": len(x), "visual_phrase_states": dict(Counter(r["visual_observation_state"] for r in phrase_rows)),
        "visual_missing_phrases": int((~mask[:,144:].all(1)).sum()),
        "pitch_missing_phrases": int((~mask[:,143]).sum()),
        "all_finite_features": bool(np.isfinite(x).all()),
        "iqr_candidate_count": len(candidates), "iqr_candidate_samples": len({r["sample_id"] for r in candidates}),
        "correlation_unit": "100 clips; separate equal-weight source-video means; no causal/significance claims",
    }
    dump(OUT / "q1_audit.json", report)
    return report


def load_pickle(path):
    with path.open("rb") as handle:
        return ArrayUnpickler(handle).load()


def arr_digest(array):
    a = np.ascontiguousarray(array)
    return hashlib.sha256(memoryview(a).cast("B")).hexdigest()


def modality_stats(x, valid, ids, detailed):
    finite = np.isfinite(x)
    zero = np.all(x == 0, axis=-1)
    result = {
        "shape": list(x.shape), "dtype": str(x.dtype), "nonfinite_values": int((~finite).sum()),
        "valid_positions": int(valid.sum()), "zero_valid_positions": int((zero & valid).sum()),
        "samples_with_zero_valid_positions": int((zero & valid).any(1).sum()),
        "entire_tensor_zero_samples": int(zero.all(1).sum()),
        "nonzero_outside_valid_positions": int((~zero & ~valid).sum()),
    }
    records = []
    if detailed:
        active = valid & ~zero & finite.all(-1)
        dims = []
        for j in range(x.shape[-1]):
            dims.append({"dimension": j, **statistics(x[:,:,j][active])})
        result["dimension_statistics"] = dims
        norm = np.sqrt(np.mean(np.square(x), axis=-1))
        for i, sid in enumerate(ids):
            values = norm[i, active[i]]
            records.append({"sample_id": sid,
                            "observed_fraction": float(active[i].sum()/max(1, valid[i].sum())),
                            "mean_position_rms": float(values.mean()) if len(values) else np.nan})
    return result, records


def q2_audit():
    labels, sheet = workbook_audit(DATA / "附件2-数据集特征文件/label.xlsx")
    label_map = {r["sample_id"]: r for r in labels}
    reports, split_ids, signatures, metrics = {}, {}, {}, {}
    for version in ("aligned_50", "unaligned_50"):
        path = DATA / "附件2-数据集特征文件" / (version+".pkl")
        data = load_pickle(path)
        report = {"file_bytes": path.stat().st_size, "sha256": sha256(path), "splits": {}}
        for split, d in data.items():
            ids = list(map(str, d["id"]))
            n = len(ids)
            bert = np.asarray(d["text_bert"])
            valid = (bert[:,1] == 1) & (bert[:,0] != 101) & (bert[:,0] != 102) & (bert[:,0] != 0)
            errors = []
            shapes = {k: list(np.shape(v)) for k,v in d.items()}
            if any(np.shape(v)[0] != n for v in d.values()):
                errors.append("field_sample_count_mismatch")
            for key, shape in {"text":(n,50,768),"text_bert":(n,3,50),
                               "audio":(n,50 if version=="aligned_50" else 500,74),
                               "vision":(n,50 if version=="aligned_50" else 500,35)}.items():
                if np.shape(d[key]) != shape: errors.append(f"{key}_shape_mismatch")
            if not (np.isfinite(bert).all() and np.all(bert == np.rint(bert))
                    and np.all((bert[:,0] >= 0) & (bert[:,0] < 30522))
                    and np.isin(bert[:,1], [0,1]).all() and (bert[:,2] == 0).all()):
                errors.append("invalid_bert_inputs")
            if np.any(np.diff(bert[:,1], axis=1) > 0):
                errors.append("attention_not_prefix")
            intensity, cls = np.asarray(d["regression_labels"]), np.asarray(d["classification_labels"])
            if not np.isfinite(intensity).all() or not np.all(np.abs(intensity) <= 3) or not np.isin(cls, [0,1,2]).all():
                errors.append("invalid_label_values")
            if not np.array_equal(np.where(intensity<0,0,np.where(intensity>0,2,1)), cls):
                errors.append("class_sign_mismatch")
            mismatches = []
            for i, sid in enumerate(ids):
                row = label_map.get(sid)
                if row is None or row.get("mode") != split or row["text"] != str(d["raw_text"][i]):
                    mismatches.append(sid)
                elif split != "test" and abs(float(row["label"]) - float(intensity[i])) > 1e-7:
                    mismatches.append(sid)
            s = {"n": n, "shapes": shapes, "errors": errors, "spreadsheet_mismatches": mismatches,
                 "duplicate_ids": duplicates(ids,ids), "modalities": {}}
            fingerprint = {k: arr_digest(np.asarray(d[k])) for k in ("raw_text","text_bert","text",
                                                                     "classification_labels","regression_labels")}
            if version == "aligned_50":
                split_ids[split] = ids
                signatures[split] = fingerprint
            else:
                s["same_ids_order_as_aligned"] = ids == split_ids[split]
                s["same_text_and_label_fields_as_aligned"] = fingerprint == signatures[split]
            features = [{"sample_id": sid, "semantic_length": int(valid[i].sum())}
                        for i,sid in enumerate(ids)]
            for modality in ("text", "audio", "vision"):
                observed_domain = valid
                if version == "unaligned_50" and modality != "text":
                    lengths = np.asarray(d[modality+"_lengths"])
                    if not np.all((lengths >= 0) & (lengths <= 500) & (lengths == np.rint(lengths))):
                        errors.append(modality+"_invalid_lengths")
                    observed_domain = np.arange(500)[None,:] < lengths[:,None]
                    s[modality+"_lengths"] = statistics(lengths)
                stats, values = modality_stats(np.asarray(d[modality]), observed_domain, ids, split != "test")
                if version == "unaligned_50" and modality != "text":
                    nonzero = np.any(d[modality] != 0, axis=-1)
                    suffix = nonzero & ~observed_domain
                    stats["samples_with_nonzero_after_declared_length"] = int(suffix.any(1).sum())
                    stats["nonzero_rows_total"] = int(nonzero.sum())
                    stats["length_equals_nonzero_row_count_samples"] = int((lengths == nonzero.sum(1)).sum())
                    first_zero = np.where(nonzero.all(1), 500, np.argmax(~nonzero,axis=1))
                    stats["length_equals_first_zero_clamped_at_one_samples"] = int(
                        (lengths == np.maximum(1,first_zero)).sum())
                    stats["declared_length_scope"] = "Statistics use declared prefix; nonzero suffix makes true valid domain uncertain."
                    stats["suffix_example_ids"] = [ids[i] for i in np.flatnonzero(suffix.any(1))[:8]]
                    if suffix.any():
                        errors.append(modality+"_nonzero_after_declared_length_requires_semantic_clarification")
                s["modalities"][modality] = stats
                for rec, val in zip(features, values):
                    rec[modality+"_observed_fraction"] = val["observed_fraction"]
                    rec[modality+"_rms"] = val["mean_position_rms"]
            if split != "test":
                s["regression_distribution"] = statistics(intensity)
                s["class_counts"] = dict(Counter(map(str, cls.astype(int))))
                s["semantic_lengths"] = statistics(valid.sum(1))
                s["duplicate_normalized_text"] = duplicates([" ".join(str(t).lower().split()) for t in d["raw_text"]], ids)
                for i, row in enumerate(features): row["label"] = float(intensity[i])
                metrics[version+"_"+split] = features
                write_csv(OUT/(version+"_"+split+"_correlations.csv"), correlation_table(features,
                          ["label","semantic_length","text_rms","audio_rms","vision_rms",
                           "audio_observed_fraction","vision_observed_fraction"]))
                write_csv(OUT/(version+"_"+split+"_samples.csv"), features)
            report["splits"][split] = s
            print(f"Attachment2 {version} {split}: {n} checked", flush=True)
        reports[version] = report
        del data, d
        gc.collect()
    report = {"spreadsheet": sheet, "versions": reports,
              "split_overlap": {a+"/"+b: {
                  "sample_ids": sorted(set(split_ids[a]) & set(split_ids[b])),
                  "video_ids": sorted({s.split("$_$")[0] for s in split_ids[a]} & {s.split("$_$")[0] for s in split_ids[b]})}
                  for a,b in itertools.combinations(split_ids,2)},
              "scope": "train/valid descriptive statistics; test schema/finite/range/sign/ID checks only, no test distribution or model evaluation"}
    q1_ids = {r["sample_id"] for r in json.loads((ROOT / "resources/manifest.json").read_text())}
    report["q1_current_overlap"] = {split:len(q1_ids & set(ids)) for split,ids in split_ids.items()}
    report["q1_ids_absent_from_attachment2"] = sorted(q1_ids-set().union(*map(set,split_ids.values())))
    dump(OUT/"q2_audit.json", report)
    return report


def special_audit():
    records, comparison = [], {}
    for attachment in ("附件3", "附件4"):
        base = next(DATA.glob(attachment+"*"))
        for path in sorted(base.rglob("*.pkl")):
            unaligned = "未对齐版本" in path.parts
            version = "unaligned" if unaligned else "aligned"
            data = load_pickle(path)
            d = data.get("test",data)
            row = {"attachment": attachment, "version": version, "file": str(path.relative_to(DATA)),
                   "sha256": sha256(path), "shapes": {k:list(np.shape(v)) for k,v in d.items()},
                   "errors": [], "notes": []}
            audio = np.asarray(d["audio"]).reshape(-1,74)
            vision = np.asarray(d["vision"]).reshape(-1,35)
            expected = 500 if unaligned else 50
            if len(audio) != expected or len(vision) != expected:
                row["errors"].append("modality_shape_mismatch")
            for name, x in (("audio",audio),("vision",vision)):
                row[name+"_nonfinite"] = int((~np.isfinite(x)).sum())
                if row[name+"_nonfinite"]: row["errors"].append(name+"_nonfinite")
            if "text_bert" in d:
                bert = np.asarray(d["text_bert"]).reshape(3,50)
                good = (np.isfinite(bert).all() and np.all(bert==np.rint(bert))
                        and np.all((bert[0]>=0)&(bert[0]<30522)) and np.isin(bert[1],[0,1]).all()
                        and (bert[2]==0).all())
                if not good: row["errors"].append("invalid_bert_inputs")
                sep = np.flatnonzero(bert[0]==102)
                if len(sep)==1:
                    semantic = (np.arange(50)>0)&(np.arange(50)<sep[0])
                    row["semantic_positions_from_sep"] = int(semantic.sum())
                    row["text_zero_holes_before_sep"] = int(((bert[0]==0)&semantic).sum())
                    row["attention_holes_before_sep"] = int(((bert[1]==0)&semantic).sum())
                else:
                    semantic = (bert[1]==1)&~np.isin(bert[0],[0,101,102])
                    row["notes"].append("semantic_extent_ambiguous_without_unique_sep")
                if not unaligned:
                    row["audio_zero_semantic_positions"] = int((np.all(audio==0,axis=1)&semantic).sum())
                    row["vision_zero_semantic_positions"] = int((np.all(vision==0,axis=1)&semantic).sum())
            else:
                row["notes"].append("no_text_bert_or_text_embeddings")
            if unaligned:
                for name,x in (("audio",audio),("vision",vision)):
                    if name+"_lengths" in d:
                        length = int(d[name+"_lengths"])
                        if not 0 <= length <= 500: row["errors"].append(name+"_length_out_of_range")
                        row[name+"_zero_within_declared_length"] = int(np.all(x[:length]==0,axis=1).sum())
                        row[name+"_nonzero_after_length"] = int(np.any(x[length:]!=0,axis=1).sum())
                        if row[name+"_nonzero_after_length"]:
                            row["notes"].append(name+"_nonzero_after_declared_length")
                    else:
                        row["notes"].append(name+"_length_missing_cannot_separate_padding")
            if "text" in d and not np.isfinite(d["text"]).all():
                row["errors"].append("text_nonfinite")
            if "raw_text" in d and not str(np.asarray(d["raw_text"]).ravel()[0]).strip():
                row["errors"].append("empty_raw_text")
            if attachment == "附件4":
                video = path.parent/"videos"/(path.stem+".mp4")
                row["video_exists"] = video.exists()
                row["id_matches_filename"] = str(d["id"]) == path.stem
                if not video.exists() or not row["id_matches_filename"]: row["errors"].append("video_pairing")
                if video.exists(): row["video_sha256"] = sha256(video)
                signature = {k:arr_digest(np.asarray(d[k])) for k in ("raw_text","text_bert","text")}
                if unaligned:
                    previous = comparison.get(path.stem)
                    row["paired_version_same_text"] = previous is not None and previous["text"] == signature
                    row["paired_version_same_video"] = previous is not None and previous["video"] == row.get("video_sha256")
                else:
                    comparison[path.stem] = {"text":signature,"video":row.get("video_sha256")}
            records.append(row)
    report = {"files": records, "counts": dict(Counter(r["attachment"]+"/"+r["version"] for r in records)),
              "files_with_errors": sum(bool(r["errors"]) for r in records),
              "interpretation": "Zero positions are observed interface states, not an estimate of the masking generator; no training settings chosen."}
    dump(OUT/"special_audit.json",report)
    return report


def main(skip_media=False, only_q1=False):
    OUT.mkdir(parents=True, exist_ok=True)
    start=time.perf_counter()
    q1=q1_audit(decode=not skip_media)
    if not only_q1:
        q2_audit()
        special_audit()
    dump(OUT/"run.json", {"elapsed_s":time.perf_counter()-start, "media_redecoded":not skip_media,
                         "only_q1":only_q1, "source_files_modified":False,
                         "quality_issue_requires_q1_rerun":bool(q1["digital_silence_ids"] or q1["visual_missing_phrases"])})
    print(f"Audit saved to reports/data_quality; {time.perf_counter()-start:.2f}s", flush=True)


if __name__ == "__main__":
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument("--skip-media",action="store_true")
    parser.add_argument("--only-q1",action="store_true")
    args=parser.parse_args()
    main(args.skip_media,args.only_q1)
